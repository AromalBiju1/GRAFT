"""Document ingestion and token chunking contract tests."""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import tiktoken
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from indexing.chunker import chunk_text
from indexing.ingest import DocumentParsingError, ingest_document, parse_document


def write_pdf(path: Path, pages: list[str]) -> None:
    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({
            NameObject('/Type'): NameObject('/Font'),
            NameObject('/Subtype'): NameObject('/Type1'),
            NameObject('/BaseFont'): NameObject('/Helvetica'),
        })
        page[NameObject('/Resources')] = DictionaryObject({
            NameObject('/Font'): DictionaryObject({NameObject('/F1'): font}),
        })
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 50 700 Td ({text}) Tj ET'.encode('ascii'))
        page[NameObject('/Contents')] = stream
    writer.write(path)


class IndexingTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.encoding = tiktoken.get_encoding('cl100k_base')

    def assert_schema(self, chunks: list[dict]) -> None:
        self.assertTrue(chunks)
        for index, chunk in enumerate(chunks):
            self.assertEqual(set(chunk), {
                'chunk_id', 'document_id', 'text', 'chunk_index', 'token_count',
            })
            self.assertIsInstance(chunk['chunk_id'], str)
            self.assertEqual(chunk['document_id'], 'test_doc')
            self.assertEqual(chunk['chunk_index'], index)
            self.assertTrue(chunk['text'].strip())
            self.assertEqual(chunk['token_count'], len(self.encoding.encode(
                chunk['text'], disallowed_special=(),
            )))

    def test_pdf_parsing_and_ingestion(self) -> None:
        path = self.root / 'sample.PDF'
        write_pdf(path, ['First readable page.', '', 'Last readable page.'])
        self.assertEqual(parse_document(path), 'First readable page.\n\nLast readable page.')
        self.assert_schema(ingest_document(path, document_id='test_doc'))

    def test_docx_parsing_and_ingestion(self) -> None:
        path = self.root / 'sample.docx'
        doc = Document()
        for paragraph in ['First paragraph.', '  ', 'Second paragraph.']:
            doc.add_paragraph(paragraph)
        doc.save(path)
        self.assertEqual(parse_document(path), 'First paragraph.\n\nSecond paragraph.')
        self.assert_schema(ingest_document(path, document_id='test_doc'))

    def test_size_overlap_order_and_no_lost_content(self) -> None:
        sentences = [f'Sentence {i} contains several unique details about this document.' for i in range(150)]
        text = ' '.join(sentences)
        chunks = chunk_text(text, document_id='test_doc')
        self.assert_schema(chunks)
        previous_end = 0
        for index, chunk in enumerate(chunks):
            start = text.index(chunk['text'])
            end = start + len(chunk['text'])
            self.assertGreater(end, previous_end)
            if index:
                self.assertLess(start, previous_end)
                overlap = text[start:previous_end]
                self.assertAlmostEqual(len(self.encoding.encode(overlap)), 50, delta=12)
            if index < len(chunks) - 1:
                self.assertGreaterEqual(chunk['token_count'], 370)
            self.assertLessEqual(chunk['token_count'], 400)
            self.assertTrue(any(chunk['text'].startswith(s) for s in sentences))
            self.assertTrue(any(chunk['text'].endswith(s) for s in sentences))
            previous_end = end
        self.assertEqual(previous_end, len(text))

    def test_paragraph_boundaries_preferred(self) -> None:
        paragraph = 'A useful sentence about document indexing. ' * 35
        text = paragraph.strip() + '\n\n' + paragraph.strip()
        chunks = chunk_text(text, overlap=0)
        self.assertEqual([c['text'] for c in chunks], [paragraph.strip()] * 2)

    def test_zero_overlap_reconstructs_all_content(self) -> None:
        sentences = [f'Item {i} has some useful content.' for i in range(100)]
        text = ' '.join(sentences)
        chunks = chunk_text(text, overlap=0, chunk_size=80)
        self.assertEqual(' '.join(c['text'] for c in chunks), text)
        self.assertTrue(all(c['token_count'] <= 80 for c in chunks))

    def test_ids_are_stable_and_document_specific(self) -> None:
        first = chunk_text('Example text.', document_id='test_doc')
        self.assertEqual(first, chunk_text('Example text.', document_id='test_doc'))
        self.assertEqual(first[0]['chunk_id'], chunk_text('Changed text.', document_id='test_doc')[0]['chunk_id'])
        self.assertNotEqual(first[0]['chunk_id'], chunk_text('Example text.', document_id='other')[0]['chunk_id'])

    def test_oversized_sentence_is_preserved_and_progresses(self) -> None:
        sentence = 'Unicode 世界 🧬 ' * 500 + 'ends.'
        chunks = chunk_text(sentence + '\n\nLast sentence.', document_id='test_doc')
        self.assert_schema(chunks)
        self.assertEqual([c['text'] for c in chunks], [sentence, 'Last sentence.'])
        self.assertGreater(chunks[0]['token_count'], 400)

    def test_literal_special_tokens_and_whitespace(self) -> None:
        chunks = chunk_text('  Literal <|endoftext|> content.\n \n ', document_id='test_doc')
        self.assert_schema(chunks)
        self.assertEqual(len(chunks), 1)

    def test_indented_paragraph_does_not_create_empty_chunk(self) -> None:
        long_sentence = 'word ' * 500 + 'ends.'
        chunks = chunk_text('Introduction.\n\n   ' + long_sentence)
        self.assertEqual([c['text'] for c in chunks], ['Introduction.', long_sentence])

    def test_empty_and_corrupt_documents(self) -> None:
        for suffix in ['.pdf', '.docx']:
            path = self.root / ('empty' + suffix)
            if suffix == '.pdf':
                write_pdf(path, [''])
            else:
                Document().save(path)
            with self.subTest(suffix=suffix, kind='no text'):
                with self.assertRaises(DocumentParsingError):
                    parse_document(path)
            for content in [b'', b'not a document']:
                path.write_bytes(content)
                with self.subTest(suffix=suffix, content=content):
                    with self.assertRaises(DocumentParsingError) as caught:
                        parse_document(path)
                    self.assertIsNotNone(caught.exception.__cause__)

    def test_unreadable_document_and_missing_path(self) -> None:
        path = self.root / 'missing.pdf'
        with self.assertRaises(DocumentParsingError) as caught:
            parse_document(path)
        self.assertIsInstance(caught.exception.__cause__, FileNotFoundError)
        write_pdf(path, ['Readable.'])
        with patch('indexing.parser.PdfReader', side_effect=PermissionError('denied')):
            with self.assertRaises(DocumentParsingError) as caught:
                parse_document(path)
        self.assertIsInstance(caught.exception.__cause__, PermissionError)

    def test_invalid_inputs(self) -> None:
        with self.assertRaisesRegex(ValueError, 'Unsupported document type'):
            parse_document(self.root / 'unsupported.txt')
        with self.assertRaises(DocumentParsingError):
            chunk_text(' \n ')
        for options in [{'chunk_size': 0}, {'overlap': -1}, {'overlap': 400}, {'document_id': ''}]:
            with self.subTest(options=options), self.assertRaises(ValueError):
                chunk_text('Some text.', **options)


if __name__ == '__main__':
    unittest.main()

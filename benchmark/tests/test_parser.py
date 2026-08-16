"""Focused tests for document text extraction."""

import tempfile
import unittest
from pathlib import Path

from docx import Document
from pypdf import PdfWriter

from indexing.parser import parse_document


class ParseDocumentTests(unittest.TestCase):
    def test_pdf_parsing_returns_readable_text(self) -> None:
        sample = (
            Path(__file__).parents[2]
            / "data"
            / "sample_docs"
            / "rfc-editor.org_rfc_rfc793.txt.pdf"
        )

        text = parse_document(sample)

        self.assertTrue(text.strip())
        self.assertIn("TRANSMISSION CONTROL PROTOCOL", text.upper())

    def test_docx_parsing_skips_empty_paragraphs(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.docx"
            document = Document()
            document.add_paragraph("First readable paragraph.")
            document.add_paragraph("   ")
            document.add_paragraph("Second readable paragraph.")
            document.save(path)

            text = parse_document(path)

        self.assertEqual(
            text,
            "First readable paragraph.\n\nSecond readable paragraph.",
        )

    def test_missing_file_raises_file_not_found_error(self) -> None:
        with self.assertRaisesRegex(FileNotFoundError, "Document not found"):
            parse_document("does-not-exist.pdf")

    def test_unsupported_extension_raises_value_error(self) -> None:
        with tempfile.NamedTemporaryFile(suffix=".txt") as file:
            with self.assertRaisesRegex(ValueError, "Unsupported document type"):
                parse_document(file.name)

    def test_pdf_without_extractable_text_raises_value_error(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "empty.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=72, height=72)
            with path.open("wb") as output:
                writer.write(output)

            with self.assertRaisesRegex(ValueError, "No extractable text"):
                parse_document(path)

    def test_docx_without_extractable_text_raises_value_error(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "empty.docx"
            Document().save(path)

            with self.assertRaisesRegex(ValueError, "No extractable text"):
                parse_document(path)


if __name__ == "__main__":
    unittest.main()
